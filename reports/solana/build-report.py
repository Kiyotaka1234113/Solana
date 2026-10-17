import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
data = json.loads((ROOT / "report-data.json").read_text(encoding="utf-8"))


def set_font(style, name, size):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for family in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{family}"), name)
    for border in style.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)


def label_paragraph(doc, label, value):
    p = doc.add_paragraph()
    p.add_run(label + ": ").bold = True
    p.add_run(value)
    return p


def code_block(doc, lines):
    p = doc.add_paragraph(style="Code")
    p.paragraph_format.keep_together = True
    p.add_run("\n".join(lines))


def hyperlink(paragraph, text, url):
    relation = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    element.append(run)
    paragraph._p.append(element)


def markdown_work(work):
    text = [f"# {work['title']}", "", f"**Статус:** {work['status']}", ""]
    for field, title in [
        ("purpose", "Цель"), ("implementation", "Реализация"),
        ("scenario", "Сценарий"), ("result", "Результат"),
    ]:
        text.extend([f"## {title}", "", work[field], ""])
    text.extend(["## Файлы", "", *[f"- `{file}`" for file in work["files"]], ""])
    if work.get("addresses"):
        text.extend(["## Адреса", ""])
        for label, address in work["addresses"]:
            text.extend([f"{label}: `{address}`", ""])
    text.extend(["## Команды", "", "```text", *work["commands"], "```", ""])
    text.extend(["## Результаты проверки", "", "```text", *work["log"], "```", ""])
    if work.get("links"):
        text.extend(["## Транзакции", "", *[f"- [{label}]({url})" for label, url in work["links"]], ""])
    text.extend(["## Итог и оставшиеся шаги", "", work["next"], ""])
    return "\n".join(text)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

for name, size in [("Normal", 11), ("Title", 23), ("Subtitle", 11), ("Heading 1", 17), ("Heading 2", 12)]:
    set_font(doc.styles[name], "Calibri", size)
doc.styles["Subtitle"].font.italic = False
for spacing in doc.styles["Subtitle"].element.xpath(".//w:rPr/w:spacing"):
    spacing.getparent().remove(spacing)
normal = doc.styles["Normal"].paragraph_format
normal.space_after = Pt(6)
normal.line_spacing = 1.08
doc.styles["Title"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(8)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
code = doc.styles.add_style("Code", 1)
set_font(code, "Consolas", 8.5)
code.paragraph_format.line_spacing = 1.05
code.paragraph_format.space_after = Pt(6)

doc.core_properties.title = data["title"]
doc.core_properties.subject = "Отчёты по шести практическим работам учебного проекта Solana"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""

doc.add_paragraph(data["title"], "Title")
doc.add_paragraph(f"Состояние проекта на {data['date']}", "Subtitle")
doc.add_paragraph(data["introduction"])
doc.add_heading("Сводка результатов", level=2)
table = doc.add_table(rows=1, cols=3)
table.autofit = False
widths = [Inches(2.0), Inches(2.4), Inches(2.6)]
for cell, width, title in zip(table.rows[0].cells, widths, ["Работа", "Содержание", "Проверка"]):
    cell.width = width
    cell.text = title
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F2F2F2")
    cell._tc.get_or_add_tcPr().append(shade)
    for run in cell.paragraphs[0].runs:
        run.bold = True
header = OxmlElement("w:tblHeader")
table.rows[0]._tr.get_or_add_trPr().append(header)
for row_data in data["summary"]:
    row = table.add_row()
    for cell, width, text in zip(row.cells, widths, row_data):
        cell.width = width
        cell.text = text
for row in table.rows:
    no_split = OxmlElement("w:cantSplit")
    row._tr.get_or_add_trPr().append(no_split)
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.line_spacing = 1.05
        margins = OxmlElement("w:tcMar")
        for edge in ("top", "left", "bottom", "right"):
            margin = OxmlElement(f"w:{edge}")
            margin.set(qn("w:w"), "95")
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        cell._tc.get_or_add_tcPr().append(margins)
doc.add_paragraph()
doc.add_paragraph(data["verification"])
doc.add_paragraph("Среда проекта: Rust, Anchor 0.30.1, TypeScript, @solana/web3.js и @solana/spl-token. Сеть для интеграционных сценариев - Devnet.")

for work in data["works"]:
    doc.add_page_break()
    doc.add_heading(work["title"], level=1)
    label_paragraph(doc, "Статус", work["status"])
    for field, title in [
        ("purpose", "Цель"), ("implementation", "Реализация"),
        ("scenario", "Сценарий проверки"), ("result", "Результат"),
    ]:
        label_paragraph(doc, title, work[field])
    label_paragraph(doc, "Файлы", "; ".join(work["files"]))
    if work.get("addresses"):
        code_block(doc, [f"{label}: {address}" for label, address in work["addresses"]])
    doc.add_heading("Команды и результаты", level=2)
    code_block(doc, work["commands"])
    code_block(doc, work["log"])
    if work.get("links"):
        p = doc.add_paragraph()
        p.add_run("Подтверждённые транзакции ").bold = True
        for i, (label, url) in enumerate(work["links"]):
            if i:
                p.add_run("; ")
            hyperlink(p, label, url)
    label_paragraph(doc, "Итог", work["next"])

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run()
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
run._r.append(field)
run.font.size = Pt(9)

docx_file = ROOT / "Solana-practical-report.docx"
doc.save(docx_file)

markdown = [f"# {data['title']}", "", f"Дата проверки: {data['date']}", "", data["introduction"], ""]
markdown.extend(["## Сводка", "", "| Работа | Содержание | Проверка |", "|---|---|---|"])
for row in data["summary"]:
    markdown.append("| " + " | ".join(row) + " |")
markdown.extend(["", data["verification"], "", "## Отдельные отчёты", ""])
for work in data["works"]:
    markdown.append(f"- [{work['title']}]({work['file']})")
markdown.append("")
for work in data["works"]:
    text = markdown_work(work)
    (ROOT / work["file"]).write_text(text, encoding="utf-8")
    markdown.extend([text, ""])
(ROOT / "README.md").write_text("\n".join(markdown).rstrip() + "\n", encoding="utf-8")
print(f"Created {docx_file.name}, README.md, and {len(data['works'])} individual reports")
